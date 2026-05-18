from io import BytesIO

import pandas as pd
from django.test import SimpleTestCase

from openpyxl import load_workbook

from .services import (
    CLASSIFICATION_BUCKETS_ORDER,
    CLASSIFICATION_SHEET_NAME,
    filter_timesheet_fact_rows,
    process_timesheet,
    process_timesheets,
)


class ProcessTimesheetTests(SimpleTestCase):
    def test_processes_multiple_files_into_single_output(self):
        source1 = pd.DataFrame(
            [
                {
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "09:00 AM",
                }
            ]
        )
        source2 = pd.DataFrame(
            [
                {
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Service Code": "TRAVL",
                    "Service Description": "Over 30 Minute Travel Time",
                    "Pay Units": 0.5,
                    "Actual Time In": "10:00 AM",
                    "Actual Time Out": "10:30 AM",
                }
            ]
        )
        files = [
            (BytesIO(source1.to_csv(index=False).encode("utf-8")), "part1.csv"),
            (BytesIO(source2.to_csv(index=False).encode("utf-8")), "part2.csv"),
        ]
        output_df, _ = process_timesheets(files)

        self.assertEqual(len(output_df), 2)
        self.assertSetEqual(set(output_df["Source File"]), {"part1.csv", "part2.csv"})
        self.assertEqual(round(float(output_df["Total Hours Worked"].sum()), 2), 1.5)
        timesheet_sheet = pd.read_excel(_, sheet_name="Timesheet")
        employee_col = set(timesheet_sheet["Employee"].astype(str))
        self.assertIn("File Grand Total (part1.csv)", employee_col)
        self.assertIn("File Grand Total (part2.csv)", employee_col)
        self.assertIn("OVERALL GRAND TOTAL", employee_col)

    def test_groups_by_employee_and_service_date(self):
        source = pd.DataFrame(
            [
                {
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Actual Time In": "08:15 AM",
                    "Actual Time Out": "09:00 AM",
                },
                {
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Actual Time In": "11:00 AM",
                    "Actual Time Out": "01:30 PM",
                },
            ]
        )
        input_stream = BytesIO()
        source.to_excel(input_stream, index=False)
        input_stream.seek(0)

        output_df, _ = process_timesheet(input_stream, filename="input.xlsx")

        self.assertEqual(len(output_df), 1)
        self.assertEqual(output_df.iloc[0]["Employee"], "Doe, Jane")
        self.assertEqual(output_df.iloc[0]["Earliest Actual Time In"], "08:15 AM")
        self.assertEqual(output_df.iloc[0]["Latest Actual Time Out"], "01:30 PM")
        self.assertEqual(output_df.iloc[0]["Total Hours Worked"], 5.25)
        self.assertEqual(output_df.iloc[0]["Total Travel Miles"], 0.0)

    def test_accepts_csv_input(self):
        source = pd.DataFrame(
            [
                {
                    "First Name": "John",
                    "Last Name": "Smith",
                    "Service Date": "2026-02-10",
                    "Actual Time In": "09:00 AM",
                    "Actual Time Out": "11:00 AM",
                }
            ]
        )
        input_stream = BytesIO(source.to_csv(index=False).encode("utf-8"))

        output_df, _ = process_timesheet(input_stream, filename="input.csv")

        self.assertEqual(len(output_df), 1)
        self.assertEqual(output_df.iloc[0]["Employee"], "Smith, John")
        self.assertEqual(output_df.iloc[0]["Total Hours Worked"], 2.0)
        self.assertEqual(output_df.iloc[0]["Total Travel Miles"], 0.0)

    def test_accepts_header_aliases(self):
        source = pd.DataFrame(
            [
                {
                    "firstname": "Mia",
                    "lastname": "Stone",
                    "DOS": "2026-03-02",
                    "Time In": "07:00",
                    "Time Out": "15:00",
                }
            ]
        )
        input_stream = BytesIO(source.to_csv(index=False).encode("utf-8"))

        output_df, _ = process_timesheet(input_stream, filename="alias.csv")

        self.assertEqual(output_df.iloc[0]["Employee"], "Stone, Mia")
        self.assertEqual(output_df.iloc[0]["Total Hours Worked"], 8.0)

    def test_accepts_tab_separated_csv_with_csv_extension(self):
        line = (
            "First Name\tLast Name\tService Date\tActual Time In\tActual Time Out\n"
            "Ann\tLee\t2026-03-10\t09:00 AM\t10:00 AM\n"
        )
        buf = BytesIO(line.encode("utf-8"))
        output_df, _ = process_timesheet(buf, filename="payroll.csv")
        self.assertEqual(output_df.iloc[0]["Employee"], "Lee, Ann")
        self.assertEqual(output_df.iloc[0]["Total Hours Worked"], 1.0)

    def test_accepts_utf8_bom_on_header(self):
        inner = "First Name,Last Name,Service Date,Actual Time In,Actual Time Out\n"
        inner += "Bob,Nguyen,2026-04-05,08:00 AM,09:30 AM\n"
        buf = BytesIO(inner.encode("utf-8-sig"))
        output_df, _ = process_timesheet(buf, filename="bom.csv")
        self.assertEqual(output_df.iloc[0]["Employee"], "Nguyen, Bob")

    def test_accepts_bom_prefix_on_first_column_name(self):
        inner = "\ufeffFirst Name,Last Name,Service Date,Actual Time In,Actual Time Out\n"
        inner += "Cara,Moss,2026-04-06,07:00 AM,08:00 AM\n"
        buf = BytesIO(inner.encode("utf-8"))
        output_df, _ = process_timesheet(buf, filename="weird.csv")
        self.assertEqual(output_df.iloc[0]["Employee"], "Moss, Cara")

    def test_groups_same_name_by_identifier(self):
        source = pd.DataFrame(
            [
                {
                    "Identifier": "A1",
                    "First Name": "Alex",
                    "Last Name": "Kim",
                    "Service Date": "2026-04-01",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "09:00 AM",
                },
                {
                    "Identifier": "B2",
                    "First Name": "Alex",
                    "Last Name": "Kim",
                    "Service Date": "2026-04-01",
                    "Actual Time In": "10:00 AM",
                    "Actual Time Out": "11:00 AM",
                },
            ]
        )
        input_stream = BytesIO()
        source.to_excel(input_stream, index=False)
        input_stream.seek(0)

        output_df, _ = process_timesheet(input_stream, filename="same_name.xlsx")

        self.assertEqual(len(output_df), 2)
        self.assertIn("Kim, Alex (A1)", set(output_df["Employee"]))
        self.assertIn("Kim, Alex (B2)", set(output_df["Employee"]))

    def test_adds_patient_call_hours_separately(self):
        source = pd.DataFrame(
            [
                {
                    "Identifier": "PV1",
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Service Code": "RNV",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "10:00 AM",
                    "Pay Units": 1.0,
                },
                {
                    "Identifier": "PV1",
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Service Code": "PatPC",
                    "Service Description": "Patient Phone Calls",
                    "Actual Time In": "05:00 PM",
                    "Actual Time Out": "05:30 PM",
                    "Pay Units": 0.5,
                },
            ]
        )
        input_stream = BytesIO(source.to_csv(index=False).encode("utf-8"))

        output_df, _ = process_timesheet(input_stream, filename="calls.csv")

        self.assertEqual(len(output_df), 1)
        self.assertEqual(output_df.iloc[0]["Employee"], "Doe, Jane (PV1)")
        self.assertEqual(output_df.iloc[0]["Earliest Actual Time In"], "08:00 AM")
        self.assertEqual(output_df.iloc[0]["Latest Actual Time Out"], "10:00 AM")
        self.assertEqual(output_df.iloc[0]["Call Hours Added"], 0.5)
        self.assertEqual(output_df.iloc[0]["Total Hours Worked"], 2.5)

    def test_adds_over_30_minute_travel_time_like_patient_calls(self):
        source = pd.DataFrame(
            [
                {
                    "Identifier": "PV1",
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Service Code": "RNV",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "10:00 AM",
                    "Pay Units": 1.0,
                },
                {
                    "Identifier": "PV1",
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Service Code": "TRAVL",
                    "Service Description": "Over 30 Minute Travel Time",
                    "Actual Time In": "05:00 PM",
                    "Actual Time Out": "05:30 PM",
                    "Pay Units": 0.4,
                },
            ]
        )
        input_stream = BytesIO(source.to_csv(index=False).encode("utf-8"))

        output_df, _ = process_timesheet(input_stream, filename="travel_time.csv")

        self.assertEqual(len(output_df), 1)
        self.assertEqual(output_df.iloc[0]["Employee"], "Doe, Jane (PV1)")
        self.assertEqual(output_df.iloc[0]["Earliest Actual Time In"], "08:00 AM")
        self.assertEqual(output_df.iloc[0]["Latest Actual Time Out"], "10:00 AM")
        self.assertEqual(output_df.iloc[0]["Call Hours Added"], 0.4)
        self.assertEqual(output_df.iloc[0]["Total Hours Worked"], 2.4)

    def test_sums_travel_miles_per_employee_day(self):
        source = pd.DataFrame(
            [
                {
                    "Identifier": "X1",
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "10:00 AM",
                    "Travel Miles": 10,
                },
                {
                    "Identifier": "X1",
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Actual Time In": "11:00 AM",
                    "Actual Time Out": "12:00 PM",
                    "Travel Miles": 5,
                },
                {
                    "Identifier": "X1",
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-02",
                    "Service Code": "PatPC",
                    "Service Description": "Patient Phone Calls",
                    "Actual Time In": "05:00 PM",
                    "Actual Time Out": "05:30 PM",
                    "Pay Units": 0.5,
                    "Travel Miles": 0,
                },
            ]
        )
        buf = BytesIO(source.to_csv(index=False).encode("utf-8"))
        buf.seek(0)
        output_df, excel_buf = process_timesheet(buf, filename="miles.csv")

        self.assertEqual(len(output_df), 1)
        self.assertEqual(output_df.iloc[0]["Total Travel Miles"], 15.0)
        # Visit span from two visits only: 8:00–12:00 (4 h); PatPC row adds 0.5 call hours → 4.5
        self.assertEqual(output_df.iloc[0]["Total Hours Worked"], 4.5)

        tms = pd.read_excel(excel_buf, sheet_name="Travel Miles Summary")
        self.assertEqual(float(tms.iloc[0]["total_miles"]), 15.0)

        kpi = pd.read_excel(excel_buf, sheet_name="Summary")
        self.assertEqual(
            float(kpi.loc[kpi["Metric"] == "Total travel miles", "Value"].iloc[0]),
            15.0,
        )

    def test_employee_hours_summary_multi_week_pivot(self):
        source = pd.DataFrame(
            [
                {
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-06",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "12:00 PM",
                },
                {
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-13",
                    "Actual Time In": "09:00 AM",
                    "Actual Time Out": "01:00 PM",
                },
            ]
        )
        input_stream = BytesIO()
        source.to_excel(input_stream, index=False)
        input_stream.seek(0)

        output_df, excel_buf = process_timesheet(input_stream, filename="input.xlsx")

        self.assertEqual(len(output_df), 2)
        detail_sum = round(float(output_df["Total Hours Worked"].sum()), 2)
        summary_df = pd.read_excel(excel_buf, sheet_name="Employee Hours Summary")
        self.assertEqual(round(float(summary_df.iloc[0]["total_hours"]), 2), detail_sum)
        # Sun–Sat weeks: 2026-01-06 (Tue) → week starting Sun 2026-01-04; 2026-01-13 → Sun 2026-01-11
        self.assertIn("2026-01-04", summary_df.columns)
        self.assertIn("2026-01-11", summary_df.columns)
        self.assertEqual(float(summary_df.iloc[0]["2026-01-04"]), 4.0)
        self.assertEqual(float(summary_df.iloc[0]["2026-01-11"]), 4.0)

        kpi = pd.read_excel(excel_buf, sheet_name="Summary")
        self.assertEqual(float(kpi.loc[kpi["Metric"] == "Total Hours", "Value"].iloc[0]), detail_sum)
        self.assertEqual(
            float(kpi.loc[kpi["Metric"] == "Sum of employee total hours", "Value"].iloc[0]),
            detail_sum,
        )

    def test_employee_hours_summary_splits_identifier_from_label(self):
        source = pd.DataFrame(
            [
                {
                    "Identifier": "PV1",
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-06",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "10:00 AM",
                },
            ]
        )
        input_stream = BytesIO()
        source.to_excel(input_stream, index=False)
        input_stream.seek(0)

        _, excel_buf = process_timesheet(input_stream, filename="input.xlsx")
        summary_df = pd.read_excel(excel_buf, sheet_name="Employee Hours Summary")
        self.assertEqual(summary_df.iloc[0]["employee_name"], "Doe, Jane")
        self.assertEqual(summary_df.iloc[0]["employee_id"], "PV1")

    def test_sunday_saturday_week_splits_monday_after_prior_saturday(self):
        """Apr 20–24 (Mon–Fri) in week Sun 4/19–Sat 4/25; Apr 27 (Mon) in next week Sun 4/26."""
        rows = []
        for d in ("2026-04-20", "2026-04-21", "2026-04-22", "2026-04-23", "2026-04-24"):
            rows.append(
                {
                    "Identifier": "NW8043",
                    "First Name": "Leah",
                    "Last Name": "Cheplick",
                    "Service Date": d,
                    "Actual Time In": "09:00 AM",
                    "Actual Time Out": "05:00 PM",
                }
            )
        rows.append(
            {
                "Identifier": "NW8043",
                "First Name": "Leah",
                "Last Name": "Cheplick",
                "Service Date": "2026-04-27",
                "Actual Time In": "09:00 AM",
                "Actual Time Out": "05:00 PM",
            }
        )
        buf = BytesIO()
        pd.DataFrame(rows).to_csv(buf, index=False)
        buf.seek(0)
        _, excel_buf = process_timesheet(buf, filename="weeks.csv")
        summary_df = pd.read_excel(excel_buf, sheet_name="Employee Hours Summary")
        self.assertIn("2026-04-19", summary_df.columns)
        self.assertIn("2026-04-26", summary_df.columns)
        self.assertEqual(float(summary_df.iloc[0]["2026-04-19"]), 40.0)
        self.assertEqual(float(summary_df.iloc[0]["2026-04-26"]), 8.0)
        self.assertEqual(float(summary_df.iloc[0]["total_hours"]), 48.0)

    def test_employee_hours_summary_uses_source_employee_column_when_present(self):
        source = pd.DataFrame(
            [
                {
                    "Employee": "Display From Payroll",
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-06",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "10:00 AM",
                },
            ]
        )
        input_stream = BytesIO(source.to_csv(index=False).encode("utf-8"))

        _, excel_buf = process_timesheet(input_stream, filename="input.csv")
        summary_df = pd.read_excel(excel_buf, sheet_name="Employee Hours Summary")
        self.assertEqual(summary_df.iloc[0]["employee_name"], "Display From Payroll")
        eid = summary_df.iloc[0]["employee_id"]
        self.assertTrue(
            eid == "" or (isinstance(eid, float) and pd.isna(eid)) or pd.isna(eid),
            msg="empty id should be blank; pandas often reads that as NaN from xlsx",
        )

    def test_associate_productivity_excel_skips_title_rows(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.cell(1, 1, "Associate Productivity")
        ws.cell(2, 1, "Revolutionary Home Health and Hospice")
        ws.cell(3, 1, "User: Jennifer Feldra")
        ws.cell(4, 1, "Agencies: … Date range: 5/3/2026 to 5/16/2026")
        ws.cell(5, 1, "Thursday, May 14, 2026")
        headers = [
            "Associate Name",
            "Week",
            "Patient Name",
            "Service Date",
            "Service",
            "Status",
            "Classification",
            "Actual In",
            "Actual Out",
        ]
        for col_idx, h in enumerate(headers, start=1):
            ws.cell(6, col_idx, h)
        ws.cell(7, 1, "Doe, Jane")
        ws.cell(7, 4, "2026-05-08")
        ws.cell(7, 5, "HHA Visit")
        ws.cell(7, 7, "Field Staff - Full Time")
        ws.cell(7, 8, "10:11")
        ws.cell(7, 9, "10:57")
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)

        output_df, excel_buf = process_timesheet(bio, filename="associate_productivity.xlsx")

        self.assertEqual(len(output_df), 1)
        self.assertEqual(output_df.iloc[0]["Employee"], "Doe, Jane")
        self.assertEqual(round(float(output_df.iloc[0]["Total Hours Worked"]), 2), 0.77)

        excel_buf.seek(0)
        wb_out = load_workbook(excel_buf, data_only=True)
        self.assertIn(CLASSIFICATION_SHEET_NAME, wb_out.sheetnames)

    def test_staff_classification_summary_sheet(self):
        source = pd.DataFrame(
            [
                {
                    "Employment Type": "Contractor",
                    "First Name": "Ann",
                    "Last Name": "Hire",
                    "Service Date": "2026-02-01",
                    "Service Description": "PT Eval",
                    "Actual Time In": "10:00 AM",
                    "Actual Time Out": "11:00 AM",
                    "Travel Miles": 3.0,
                },
                {
                    "Employment Type": "Contractor",
                    "First Name": "Ann",
                    "Last Name": "Hire",
                    "Service Date": "2026-02-01",
                    "Service Description": "PT Eval",
                    "Actual Time In": "02:00 PM",
                    "Actual Time Out": "03:00 PM",
                    "Travel Miles": 2.0,
                },
                {
                    "Employment Type": "Field Staff - Part Time",
                    "First Name": "Ben",
                    "Last Name": "Part",
                    "Service Date": "2026-02-01",
                    "Service Description": "RN SOC Assessment",
                    "Actual Time In": "09:00 AM",
                    "Actual Time Out": "01:00 PM",
                    "Travel Miles": 10.0,
                },
                {
                    "Employment Type": "Field Staff - Part Time",
                    "First Name": "Ben",
                    "Last Name": "Part",
                    "Service Date": "2026-02-01",
                    "Service Code": "PatPC",
                    "Service Description": "Patient Phone Calls",
                    "Actual Time In": "05:00 PM",
                    "Actual Time Out": "05:30 PM",
                    "Pay Units": 0.25,
                    "Travel Miles": 0.0,
                },
                {
                    "Employment Type": "Full time",
                    "First Name": "Carla",
                    "Last Name": "Full",
                    "Service Date": "2026-02-01",
                    "Service Description": "HHA Visit",
                    "Actual Time In": "01:00 PM",
                    "Actual Time Out": "02:00 PM",
                    "Travel Miles": 1.0,
                },
            ]
        )
        buf = BytesIO(source.to_csv(index=False).encode("utf-8"))
        buf.seek(0)
        _, excel_buf = process_timesheet(buf, filename="classified.csv")

        excel_buf.seek(0)
        wb = load_workbook(excel_buf, data_only=True)
        self.assertIn(CLASSIFICATION_SHEET_NAME, wb.sheetnames)
        ws = wb[CLASSIFICATION_SHEET_NAME]
        rows = [tuple(c.value for c in r) for r in ws.iter_rows()]

        def first_total_miles_after_title(title: str) -> float:
            seen = False
            for row in rows:
                v0 = row[0] if row else None
                if v0 == title:
                    seen = True
                elif seen and v0 == "Total miles (section)":
                    return float(row[1])
            raise AssertionError(f"no Total miles after {title!r}")

        self.assertEqual(first_total_miles_after_title("Contractors"), 5.0)
        self.assertEqual(first_total_miles_after_title("Field Staff - Part Time"), 10.0)

        found_ann_pt_eval = False
        found_ben_rn = False
        in_contractors = False
        in_pt = False
        for row in rows:
            v0 = row[0] if row else None
            if v0 == "Contractors":
                in_contractors = True
                in_pt = False
            elif v0 == "Field Staff - Part Time":
                in_pt = True
                in_contractors = False
            elif v0 in CLASSIFICATION_BUCKETS_ORDER:
                in_pt = False
                in_contractors = False
            if in_contractors and len(row) >= 5 and row[1] == "Hire, Ann" and row[2] == "pt eval":
                self.assertEqual(int(row[3]), 2)
                found_ann_pt_eval = True
            if in_pt and len(row) >= 5 and row[1] == "Part, Ben" and row[2] == "rn soc assessment":
                self.assertEqual(int(row[3]), 1)
                self.assertEqual(float(row[4]), 4.0)
                found_ben_rn = True
        self.assertTrue(found_ann_pt_eval)
        self.assertTrue(found_ben_rn)

    def test_ft_iv_visits_by_week_on_classification_sheet(self):
        source = pd.DataFrame(
            [
                {
                    "Employment Type": "Field Staff - Full Time",
                    "First Name": "Dana",
                    "Last Name": "Ivory",
                    "Week": "2026-02-02 - 2026-02-08",
                    "Service Date": "2026-02-03",
                    "Service Description": "IV Visit",
                    "Actual Time In": "09:00 AM",
                    "Actual Time Out": "10:30 AM",
                    "Travel Miles": 0,
                },
                {
                    "Employment Type": "Field Staff - Full Time",
                    "First Name": "Dana",
                    "Last Name": "Ivory",
                    "Week": "2026-02-02 - 2026-02-08",
                    "Service Date": "2026-02-04",
                    "Service Description": "HHA Visit",
                    "Actual Time In": "11:00 AM",
                    "Actual Time Out": "12:00 PM",
                    "Travel Miles": 0,
                },
                {
                    "Employment Type": "Field Staff - Full Time",
                    "First Name": "Dana",
                    "Last Name": "Ivory",
                    "Week": "2026-02-09 - 2026-02-15",
                    "Service Date": "2026-02-10",
                    "Service Description": "IV Start",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "09:00 AM",
                    "Travel Miles": 0,
                },
            ]
        )
        buf = BytesIO(source.to_csv(index=False).encode("utf-8"))
        buf.seek(0)
        _, excel_buf = process_timesheet(buf, filename="ft_iv.csv")
        excel_buf.seek(0)
        wb = load_workbook(excel_buf, data_only=True)
        ws = wb[CLASSIFICATION_SHEET_NAME]
        rows = [tuple(c.value for c in r) for r in ws.iter_rows()]

        in_ft_iv = False
        week1_hours = None
        week2_hours = None
        for row in rows:
            v0 = row[0] if row else None
            if v0 == "IV visits by week (full-time staff)":
                in_ft_iv = True
                continue
            if in_ft_iv and v0 in CLASSIFICATION_BUCKETS_ORDER:
                break
            if in_ft_iv and len(row) >= 5 and row[1] == "Ivory, Dana":
                if row[2] == "2026-02-02 - 2026-02-08":
                    week1_hours = float(row[4])
                elif row[2] == "2026-02-09 - 2026-02-15":
                    week2_hours = float(row[4])
        self.assertAlmostEqual(week1_hours, 1.5, places=2)
        self.assertAlmostEqual(week2_hours, 1.0, places=2)

    def test_filter_timesheet_fact_rows_excludes_totals(self):
        frame = pd.DataFrame(
            {
                "Employee": ["Doe, Jane", "Grand Total (Doe, Jane)", "OVERALL GRAND TOTAL"],
                "Total Hours Worked": [3.0, 3.0, 3.0],
            }
        )
        filtered = filter_timesheet_fact_rows(frame)
        self.assertEqual(len(filtered), 1)
        self.assertEqual(filtered.iloc[0]["Employee"], "Doe, Jane")

    def test_large_export_style_rounding_drift_does_not_fail_validation(self):
        source = pd.DataFrame(
            [
                {
                    "First Name": "Jane",
                    "Last Name": "Doe",
                    "Service Date": "2026-01-06",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "08:01 AM",
                },
                {
                    "First Name": "John",
                    "Last Name": "Smith",
                    "Service Date": "2026-01-06",
                    "Actual Time In": "08:00 AM",
                    "Actual Time Out": "08:01 AM",
                },
            ]
        )
        input_stream = BytesIO(source.to_csv(index=False).encode("utf-8"))

        output_df, excel_buf = process_timesheet(input_stream, filename="rounding.csv")

        self.assertEqual(len(output_df), 2)
        detail_total = round(float(output_df["Total Hours Worked"].sum()), 2)
        summary_df = pd.read_excel(excel_buf, sheet_name="Summary")
        summary_total = float(summary_df.loc[summary_df["Metric"] == "Sum of employee total hours", "Value"].iloc[0])
        self.assertEqual(summary_total, detail_total)
